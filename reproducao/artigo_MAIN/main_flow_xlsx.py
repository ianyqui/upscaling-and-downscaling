import pickle, json, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as L
F = {tuple(k.split("|")): {kk: np.array(vv) for kk, vv in v.items()} for k, v in pickle.load(open("main_flow_fields.pkl", "rb")).items()}
d = json.load(open("main_flow.json"))
PR = [("1.1", "P1"), ("1.3", "P2"), ("1.4", "P3")]
V = [("UPS", "Sem compressão"), ("DA25", "DA 25 indep."), ("DAI", "DA 1000 indep."), ("DAC", "DA 1000 cópula")]
f = Font(name="Arial", size=10); fb = Font(name="Arial", size=10, bold=True); hdr = PatternFill("solid", fgColor="DCE6F1")
thin = Border(bottom=Side(style="thin", color="999999"))
wb = Workbook(); ws = wb.active; ws.title = "Leia-me"
txt = ["Resultados numéricos: micromalha (15 × 15) × macromalha (5 × 5) — artigo MAIN, Seção 5.7",
 "",
 "Referência (Fina): média e desvio-padrão da pressão das 1000 soluções na micromalha, com média nas 3 × 3 células finas de cada célula grossa.",
 "Ensembles grossos (macromalha, 5 × 5, TPFA com tensor diagonal):",
 "  Sem compressão: tensores equivalentes de cada uma das 1000 realizações (isola o erro do upscaling).",
 "  DA 25 indep.: execução de produção, 25 realizações grossas amostradas célula a célula de forma independente.",
 "  DA 1000 indep.: 1000 realizações amostradas da mesma forma.",
 "  DA 1000 cópula: 1000 realizações amostradas dos mesmos valores representativos com cópula gaussiana (Kxx e Kyy de todas as células).",
 "Erro absoluto = grosso − fino (unidades da queda de pressão imposta, que vale 1). Erro relativo = |erro| / |fino|.",
 "Macrocélula (i, j), como nos eixos das Figuras 7 e 8: i = índice na direção x (i = 0 à esquerda), j = índice na direção y (j = 0 embaixo); centro em x = 3i + 1,5 m, y = 3j + 1,5 m.",
 "Problema 1: u = 1 à esquerda, u = 0 à direita, sem fluxo em cima e embaixo. Problema 2: u = 1 embaixo, u = 0 à esquerda, sem fluxo à direita e em cima. Problema 3: u = 1 à esquerda, u = 0 nas outras três faces.",
 "Os valores dos campos (colunas azuis) vêm de main_flow.py (arquivo main_flow_fields.pkl); os erros e o resumo são fórmulas e se recalculam.",
 "O piso amostral (duas metades de 500 realizações finas comparadas entre si) está na aba Resumo, copiado de main_flow.json."]
for i, t in enumerate(txt, 1):
    ws.cell(i, 1, t).font = fb if i == 1 else f
ws.column_dimensions["A"].width = 150
blue = Font(name="Arial", size=10, color="0000FF")
sheets = {}
for p, pn in PR:
    for stat, sn in (("mean", "média"), ("sd", "desvio")):
        sh = wb.create_sheet(f"{pn} {sn}"); sheets[(p, stat)] = sh
        heads = ["i", "j", "x (m)", "y (m)", "Fina"] + [n for _, n in V] + [f"Erro abs. {n}" for _, n in V] + [f"Erro rel. {n}" for _, n in V]
        for c, h in enumerate(heads, 1):
            cl = sh.cell(1, c, h); cl.font = fb; cl.fill = hdr; cl.alignment = Alignment(wrap_text=True, horizontal="center"); cl.border = thin
        sh.row_dimensions[1].height = 42
        r = 2
        for I in range(5):
            for J in range(5):
                sh.cell(r, 1, J); sh.cell(r, 2, I); sh.cell(r, 3, 3*J+1.5); sh.cell(r, 4, 3*I+1.5)
                c = sh.cell(r, 5, float(F[(p, "REF")][stat][I, J])); c.font = blue; c.number_format = "0.00000"
                for k, (key, _) in enumerate(V):
                    c = sh.cell(r, 6+k, float(F[(p, key)][stat][I, J])); c.font = blue; c.number_format = "0.00000"
                    e = sh.cell(r, 10+k, f"={L(6+k)}{r}-$E{r}"); e.number_format = "0.00000"; e.font = f
                    q = sh.cell(r, 14+k, f"=ABS({L(10+k)}{r})/ABS($E{r})"); q.number_format = "0.0%"; q.font = f
                for c in range(1, 5): sh.cell(r, c).font = f
                r += 1
        for c in range(1, 18): sh.column_dimensions[L(c)].width = 7 if c <= 4 else 13
        sh.freeze_panes = "E2"
# summary with formulas
sm = wb.create_sheet("Resumo", 1)
heads = ["Problema", "Estatística", "Ensemble grosso", "Erro abs. máx.", "Erro abs. médio", "Erro rel. máx.", "Erro rel. médio", "Erro L2 relativo", "R2"]
for c, h in enumerate(heads, 1):
    cl = sm.cell(1, c, h); cl.font = fb; cl.fill = hdr; cl.border = thin; cl.alignment = Alignment(wrap_text=True, horizontal="center")
r = 2
for p, pn in PR:
    for stat, sn in (("mean", "média"), ("sd", "desvio")):
        sh = sheets[(p, stat)]; nm = f"'{sh.title}'"
        for k, (key, n) in enumerate(V):
            ea = f"{nm}!{L(10+k)}2:{L(10+k)}26"; er = f"{nm}!{L(14+k)}2:{L(14+k)}26"; rf = f"{nm}!$E$2:$E$26"
            vals = [pn, sn, n, f"=MAX(MAX({ea}),-MIN({ea}))", f"=SUMPRODUCT(ABS({ea}))/25", f"=MAX({er})", f"=AVERAGE({er})",
                    f"=SQRT(SUMSQ({ea}))/SQRT(SUMSQ({rf}))", f"=1-SUMSQ({ea})/DEVSQ({rf})"]
            for c, v in enumerate(vals, 1):
                cl = sm.cell(r, c, v); cl.font = f
                cl.number_format = {4: "0.00000", 5: "0.00000", 6: "0.0%", 7: "0.0%", 8: "0.00%", 9: "0.0000"}.get(c, "General")
            r += 1
        nf = d[p]["noise_floor"][stat]
        vals = [pn, sn, "Piso: fina × fina (500/500)", nf["max_abs"], nf["mean_abs"], nf["max_rel"]/100, nf["mean_rel"]/100, nf["rel_L2"]/100, nf["R2"]]
        for c, v in enumerate(vals, 1):
            cl = sm.cell(r, c, v); cl.font = Font(name="Arial", size=10, color="0000FF", italic=True)
            cl.number_format = {4: "0.00000", 5: "0.00000", 6: "0.0%", 7: "0.0%", 8: "0.00%", 9: "0.0000"}.get(c, "General")
        r += 1
sm.cell(r+1, 1, "Linhas 'Piso' em azul itálico: valores fixos copiados de main_flow.json (comparação entre duas metades independentes do ensemble fino); demais linhas são fórmulas sobre as abas por problema.").font = f
for c, w in enumerate([9, 11, 26, 14, 14, 13, 13, 14, 9], 1): sm.column_dimensions[L(c)].width = w
sm.freeze_panes = "A2"
wb.save("Resultados_Micro_x_Macro_MAIN.xlsx"); print("saved")
